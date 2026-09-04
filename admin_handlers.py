"""
معالجات الأدمن الرئيسية
يتضمن: القائمة، الصيانة، الإحصائيات (+ زر التقييمات الجديد)، تصفير، تصدير Excel احترافي
"""
import io

from telegram import Update, InlineKeyboardButton, InlineKeyboardMarkup, InputFile
from telegram.ext import CallbackContext
from telegram.constants import ParseMode
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, PieChart, Reference

import openpyxl

from config import logger
from data_manager import get_data_manager
from bot_ui import show_main_menu
from keyboards import build_back_to_admin_menu_keyboard
from utils import is_admin_in_group

# --- ثوابت التنسيق ---
HEADER_FONT = Font(name='Arial', size=12, bold=True, color='FFFFFF')
HEADER_FILL = PatternFill(start_color='2F5496', end_color='2F5496', fill_type='solid')
HEADER_ALIGNMENT = Alignment(horizontal='center', vertical='center', wrap_text=True)
TITLE_FONT = Font(name='Arial', size=14, bold=True, color='2F5496')
TITLE_ALIGNMENT = Alignment(horizontal='center', vertical='center')
DATA_FONT = Font(name='Arial', size=11)
DATA_ALIGNMENT = Alignment(horizontal='right', vertical='center', wrap_text=True)
MONEY_FORMAT = '#,##0'
THIN_BORDER = Border(
    left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))
ALT_ROW_FILL = PatternFill(start_color='D6E4F0', end_color='D6E4F0', fill_type='solid')
SUCCESS_FILL = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
WARNING_FILL = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
DANGER_FILL = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')


# --- المعالجات ---

async def admin_main_menu_callback(update: Update, context: CallbackContext) -> None:
    await show_main_menu(update, context)


async def admin_toggle_maintenance(update: Update, context: CallbackContext) -> None:
    query = update.callback_query
    await query.answer()

    dm = get_data_manager()
    new_status = await dm.toggle_maintenance_mode()
    status = "🔧 مفعل" if new_status else "✅ غير مفعل"

    await query.edit_message_text(f"🔧 تم تغيير وضع الصيانة إلى: {status}")
    await show_main_menu(update, context)


async def admin_stats_handler(update: Update, context: CallbackContext) -> None:
    """الإحصائيات الأساسية (+ زر إحصائيات التقييمات - محسّن)"""
    query = update.callback_query
    await query.answer()

    dm = get_data_manager()
    stats = await dm.get_stats()
    new_users_week = await dm.get_new_users_count(days=7)

    stats_text = (
        f"📊 <b>إحصائيات البوت</b>\n"
        f"👥 <b>إجمالي المستخدمين النشطين:</b> {stats['total_users']}"
        f"  (+{new_users_week} هذا الأسبوع)\n\n"
        f"📋 <b>حالة الطلبات</b>\n"
        f"⏳ <b>قيد المراجعة:</b> {stats['pending']}\n"
        f"✅ <b>مفتوحة:</b> {stats['approved']}\n"
        f"✅ <b>مكتملة:</b> {stats['completed']}\n"
        f"❌ <b>مرفوضة:</b> {stats['rejected']}\n\n"
        f"💰 <b>الإحصائيات المالية</b>\n"
        f"💵 <b>إجمالي الإيرادات (من الطلبات المكتملة):</b> {stats['revenue']:,.0f} ل.س"
    )

    keyboard = [
        [InlineKeyboardButton("📈 إحصائيات متقدمة", callback_data="admin_advanced_stats")],
        [InlineKeyboardButton("⭐ إحصائيات التقييمات", callback_data="admin_ratings_stats")],
        [InlineKeyboardButton("⚠️ تصفير الإحصائيات", callback_data="admin_reset_stats")],
        [InlineKeyboardButton("🔙 العودة للقائمة الرئيسية", callback_data="admin_main_menu")]
    ]

    await query.edit_message_text(
        stats_text, reply_markup=InlineKeyboardMarkup(keyboard),
        parse_mode=ParseMode.HTML)


async def admin_advanced_stats_handler(update: Update, context: CallbackContext) -> None:
    """الإحصائيات المتقدمة (+ الأقسام الأكثر طلباً - جديد)"""
    query = update.callback_query
    await query.answer()

    dm = get_data_manager()
    daily_revenue = await dm.get_daily_revenue(days=7)
    popular_services = await dm.get_popular_services(limit=5)
    popular_specs = await dm.get_popular_specs(limit=5)
    popular_courses = await dm.get_popular_courses(limit=5)

    stats_text = "📈 <b>الإحصائيات المتقدمة</b>\n\n"

    stats_text += "📅 <b>الإيرادات اليومية (آخر 7 أيام):</b>\n"
    if daily_revenue:
        for day in daily_revenue:
            stats_text += f"  {day['day']}: {day['revenue']:,.0f} ل.س ({day['orders_count']} طلب)\n"
    else:
        stats_text += "  لا توجد بيانات كافية.\n"

    stats_text += "\n🏆 <b>الخدمات الأكثر طلباً:</b>\n"
    if popular_services:
        for i, svc in enumerate(popular_services, 1):
            stats_text += f"  {i}. {svc['service']}: {svc['count']} طلب ({svc['revenue']:,.0f} ل.س)\n"
    else:
        stats_text += "  لا توجد بيانات كافية.\n"

    stats_text += "\n🎓 <b>التخصصات الأكثر طلباً:</b>\n"
    if popular_specs:
        for i, spc in enumerate(popular_specs, 1):
            stats_text += f"  {i}. {spc['spec']}: {spc['count']} طلب\n"
    else:
        stats_text += "  لا توجد بيانات كافية.\n"

    stats_text += "\n📚 <b>المواد الأكثر طلباً:</b>\n"
    if popular_courses:
        for i, crs in enumerate(popular_courses, 1):
            stats_text += f"  {i}. {crs['course']}: {crs['count']} طلب\n"
    else:
        stats_text += "  لا توجد بيانات كافية.\n"

    keyboard = [
        [InlineKeyboardButton("📊 العودة للإحصائيات الأساسية", callback_data="admin_stats")],
        [InlineKeyboardButton("🔙 العودة للقائمة الرئيسية", callback_data="admin_main_menu")]
    ]

    await query.edit_message_text(
        stats_text, reply_markup=InlineKeyboardMarkup(keyboard),
        parse_mode=ParseMode.HTML)


async def admin_reset_stats_handler(update: Update, context: CallbackContext) -> None:
    query = update.callback_query
    await query.answer()

    keyboard = [
        [InlineKeyboardButton("✅ نعم، تصفير", callback_data="confirm_reset_stats")],
        [InlineKeyboardButton("❌ إلغاء", callback_data="admin_stats")]
    ]

    await query.edit_message_text(
        "⚠️ <b>تنبيه هام!</b>\n\n"
        "هل أنت متأكد من تصفير الإحصائيات؟\n\n"
        "🗑 سيتم حذف الطلبات المكتملة والمرفوضة نهائياً.\n"
        "⚠️ هذا الإجراء لا يمكن التراجع عنه!\n\n"
        "<b>سيتم الاحتفاظ بـ:</b>\n"
        "- ⏳ الطلبات قيد المراجعة\n"
        "- ✅ الطلبات الموافق عليها (المفتوحة)",
        reply_markup=InlineKeyboardMarkup(keyboard),
        parse_mode=ParseMode.HTML)


async def admin_confirm_reset_stats(update: Update, context: CallbackContext) -> None:
    query = update.callback_query
    await query.answer()

    dm = get_data_manager()
    deleted_count = await dm.delete_orders_by_status(['مكتمل', 'مرفوض'])
    stats = await dm.get_stats()

    result_text = (
        f"✅ <b>تم تصفير الإحصائيات بنجاح!</b>\n\n"
        f"🗑 <b>عدد الطلبات المحذوفة:</b> {deleted_count}\n\n"
        f"<b>📊 الإحصائيات الجديدة:</b>\n"
        f"- ⏳ قيد المراجعة: {stats['pending']}\n"
        f"- ✅ مفتوحة: {stats['approved']}\n"
        f"- ✅ مكتملة: {stats['completed']}\n"
        f"- ❌ مرفوضة: {stats['rejected']}\n"
        f"- 💵 إجمالي الإيرادات: {stats['revenue']:,.0f} ل.س"
    )

    await query.edit_message_text(
        result_text, reply_markup=build_back_to_admin_menu_keyboard(),
        parse_mode=ParseMode.HTML)


# --- دوال تنسيق Excel ---

def _apply_header_style(ws, row_num: int, max_col: int):
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER


def _apply_data_style(ws, start_row: int, end_row: int, max_col: int, money_cols: list = None):
    for row in range(start_row, end_row + 1):
        for col in range(1, max_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = DATA_FONT
            cell.alignment = DATA_ALIGNMENT
            cell.border = THIN_BORDER
            if (row - start_row) % 2 == 1:
                cell.fill = ALT_ROW_FILL
            if money_cols and col in money_cols:
                cell.number_format = MONEY_FORMAT


def _auto_adjust_columns(ws):
    for col in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_length + 4, 40)


def _add_title_row(ws, title: str, max_col: int, row: int = 1):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max_col)
    title_cell = ws.cell(row=row, column=1, value=title)
    title_cell.font = TITLE_FONT
    title_cell.alignment = TITLE_ALIGNMENT


# --- تصدير Excel ---

async def export_data_command(update: Update, context: CallbackContext) -> None:
    """تصدير Excel احترافي (6 أوراق مع رسوم بيانية + عمود القسم)"""
    if not await is_admin_in_group(update.effective_user.id, context):
        await update.message.reply_text("⛔ ليس لديك الصلاحية.")
        return

    await update.message.reply_text("📊 جاري تحضير ملف التصدير الاحترافي...")

    dm = get_data_manager()
    wb = openpyxl.Workbook()

    # ═══ ورقة 1: لوحة المعلومات ═══
    ws_dash = wb.active
    ws_dash.title = "لوحة المعلومات"
    ws_dash.sheet_view.rightToLeft = True

    stats = await dm.get_stats()
    new_users_week = await dm.get_new_users_count(days=7)
    new_users_month = await dm.get_new_users_count(days=30)

    _add_title_row(ws_dash, "لوحة المعلومات - بوت علوم الإدارة وإدارة الأعمال", 2, 1)
    ws_dash.row_dimensions[1].height = 35

    dashboard_data = [
        ["البند", "القيمة"],
        ["إجمالي المستخدمين", stats['total_users']],
        ["مستخدمون جدد (هذا الأسبوع)", new_users_week],
        ["مستخدمون جدد (هذا الشهر)", new_users_month],
        ["إجمالي الطلبات", stats['total_orders']],
        ["طلبات قيد المراجعة", stats['pending']],
        ["طلبات مفتوحة", stats['approved']],
        ["طلبات مكتملة", stats['completed']],
        ["طلبات مرفوضة", stats['rejected']],
        ["إجمالي الإيرادات (ل.س)", stats['revenue']],
    ]

    for i, row_data in enumerate(dashboard_data):
        for j, value in enumerate(row_data):
            ws_dash.cell(row=i + 3, column=j + 1, value=value)

    _apply_header_style(ws_dash, 3, 2)
    _apply_data_style(ws_dash, 4, 4 + len(dashboard_data) - 2, 2, money_cols=[2])
    revenue_row = 4 + len(dashboard_data) - 2
    for col in range(1, 3):
        ws_dash.cell(row=revenue_row, column=col).font = Font(
            name='Arial', size=11, bold=True, color='006100')
        ws_dash.cell(row=revenue_row, column=col).fill = SUCCESS_FILL
    _auto_adjust_columns(ws_dash)

    # ═══ ورقة 2: الطلبات ═══
    ws_orders = wb.create_sheet("الطلبات")
    ws_orders.sheet_view.rightToLeft = True

    _add_title_row(ws_orders, "سجل الطلبات", 11, 1)
    ws_orders.row_dimensions[1].height = 35

    orders_headers = [
        "رقم الطلب", "اسم الطالب", "المستخدم", "التخصص", "الخدمة",
        "المواد", "طريقة الدفع", "السعر (ل.س)", "الحالة",
        "تاريخ الإنشاء", "آخر تحديث"
    ]
    for j, header in enumerate(orders_headers):
        ws_orders.cell(row=3, column=j + 1, value=header)
    _apply_header_style(ws_orders, 3, len(orders_headers))

    orders_data = []
    for key, value in await dm.get_all_orders():
        order_copy = value.copy()
        order_copy['order_key'] = key
        courses = order_copy.get('courses', [])
        order_copy['courses'] = ', '.join(courses) if isinstance(courses, list) else str(courses)
        orders_data.append(order_copy)

    STATUS_FILLS = {
        'قيد المراجعة': WARNING_FILL, 'موافق عليه': SUCCESS_FILL,
        'مكتمل': SUCCESS_FILL, 'مرفوض': DANGER_FILL,
    }

    for i, order in enumerate(orders_data):
        row = i + 4
        values = [
            order.get('order_key', ''), order.get('full_name', ''),
            f"@{order.get('username', '')}" if order.get('username') else '',
            order.get('specialization', '') or '', order.get('service', ''),
            order.get('courses', ''), order.get('payment_method', ''),
            order.get('total_price', 0), order.get('status', ''),
            order.get('created_at', ''), order.get('updated_at', ''),
        ]
        for j, val in enumerate(values):
            ws_orders.cell(row=row, column=j + 1, value=val)

    if orders_data:
        end_row = 3 + len(orders_data)
        _apply_data_style(ws_orders, 4, end_row, len(orders_headers), money_cols=[8])
        for i in range(len(orders_data)):
            status_cell = ws_orders.cell(row=i + 4, column=9)
            if status_cell.value in STATUS_FILLS:
                status_cell.fill = STATUS_FILLS[status_cell.value]
    _auto_adjust_columns(ws_orders)

    # ═══ ورقة 3: المستخدمين ═══
    ws_users = wb.create_sheet("المستخدمين")
    ws_users.sheet_view.rightToLeft = True

    _add_title_row(ws_users, "سجل المستخدمين", 5, 1)
    ws_users.row_dimensions[1].height = 35

    users_headers = ["معرف المستخدم", "اسم المستخدم", "الاسم", "محظور", "تاريخ التسجيل"]
    for j, header in enumerate(users_headers):
        ws_users.cell(row=3, column=j + 1, value=header)
    _apply_header_style(ws_users, 3, len(users_headers))

    users_data = await dm.get_all_users()
    banned_users = await dm.get_banned_users()

    for i, (uid, uinfo) in enumerate(users_data.items()):
        row = i + 4
        ws_users.cell(row=row, column=1, value=uid)
        ws_users.cell(row=row, column=2, value=uinfo.get('username', '') or '')
        ws_users.cell(row=row, column=3, value=uinfo.get('first_name', ''))
        ws_users.cell(row=row, column=4, value='نعم' if uid in banned_users else 'لا')
        if uid in banned_users:
            ws_users.cell(row=row, column=4).fill = DANGER_FILL
            ws_users.cell(row=row, column=4).font = Font(
                name='Arial', size=11, bold=True, color='9C0006')

    if users_data:
        _apply_data_style(ws_users, 4, 3 + len(users_data), len(users_headers))
    _auto_adjust_columns(ws_users)

    # ═══ ورقة 4: الإيرادات اليومية + رسم بياني ═══
    ws_daily = wb.create_sheet("الإيرادات اليومية")
    ws_daily.sheet_view.rightToLeft = True
    _add_title_row(ws_daily, "الإيرادات اليومية (آخر 30 يوم)", 3, 1)
    ws_daily.row_dimensions[1].height = 35

    daily_headers = ["التاريخ", "عدد الطلبات", "الإيرادات (ل.س)"]
    for j, header in enumerate(daily_headers):
        ws_daily.cell(row=3, column=j + 1, value=header)
    _apply_header_style(ws_daily, 3, len(daily_headers))

    daily_revenue = await dm.get_daily_revenue(days=30)
    for i, day in enumerate(daily_revenue):
        row = i + 4
        ws_daily.cell(row=row, column=1, value=day['day'])
        ws_daily.cell(row=row, column=2, value=day['orders_count'])
        ws_daily.cell(row=row, column=3, value=day['revenue'])

    if daily_revenue:
        end_row = 3 + len(daily_revenue)
        _apply_data_style(ws_daily, 4, end_row, len(daily_headers), money_cols=[3])
        if len(daily_revenue) > 1:
            chart = BarChart()
            chart.title = "الإيرادات اليومية"
            chart.y_axis.title = "ل.س"
            chart.x_axis.title = "التاريخ"
            chart.style = 10
            chart.width = 20
            chart.height = 12
            data_ref = Reference(ws_daily, min_col=3, min_row=3, max_row=end_row)
            cats_ref = Reference(ws_daily, min_col=1, min_row=4, max_row=end_row)
            chart.add_data(data_ref, titles_from_data=True)
            chart.set_categories(cats_ref)
            ws_daily.add_chart(chart, f"A{end_row + 3}")
    _auto_adjust_columns(ws_daily)

    # ═══ ورقة 5: الخدمات الشائعة + رسم دائري ═══
    ws_services = wb.create_sheet("الخدمات الشائعة")
    ws_services.sheet_view.rightToLeft = True
    _add_title_row(ws_services, "الخدمات الأكثر طلباً", 3, 1)
    ws_services.row_dimensions[1].height = 35

    services_headers = ["الخدمة", "عدد الطلبات", "الإيرادات (ل.س)"]
    for j, header in enumerate(services_headers):
        ws_services.cell(row=3, column=j + 1, value=header)
    _apply_header_style(ws_services, 3, len(services_headers))

    popular_services = await dm.get_popular_services(limit=10)
    for i, svc in enumerate(popular_services):
        row = i + 4
        ws_services.cell(row=row, column=1, value=svc['service'])
        ws_services.cell(row=row, column=2, value=svc['count'])
        ws_services.cell(row=row, column=3, value=svc['revenue'])

    if popular_services:
        end_row = 3 + len(popular_services)
        _apply_data_style(ws_services, 4, end_row, len(services_headers), money_cols=[3])
        if len(popular_services) > 1:
            pie = PieChart()
            pie.title = "توزيع الطلبات حسب الخدمة"
            pie.style = 10
            pie.width = 18
            pie.height = 12
            data_ref = Reference(ws_services, min_col=2, min_row=3, max_row=end_row)
            cats_ref = Reference(ws_services, min_col=1, min_row=4, max_row=end_row)
            pie.add_data(data_ref, titles_from_data=True)
            pie.set_categories(cats_ref)
            ws_services.add_chart(pie, f"A{end_row + 3}")
    _auto_adjust_columns(ws_services)

    # ═══ ورقة 6: المواد الشائعة + رسم أفقي ═══
    ws_courses = wb.create_sheet("المواد الشائعة")
    ws_courses.sheet_view.rightToLeft = True
    _add_title_row(ws_courses, "المواد الأكثر طلباً", 2, 1)
    ws_courses.row_dimensions[1].height = 35

    courses_headers = ["المادة", "عدد الطلبات"]
    for j, header in enumerate(courses_headers):
        ws_courses.cell(row=3, column=j + 1, value=header)
    _apply_header_style(ws_courses, 3, len(courses_headers))

    popular_courses = await dm.get_popular_courses(limit=20)
    for i, crs in enumerate(popular_courses):
        row = i + 4
        ws_courses.cell(row=row, column=1, value=crs['course'])
        ws_courses.cell(row=row, column=2, value=crs['count'])

    if popular_courses:
        end_row = 3 + len(popular_courses)
        _apply_data_style(ws_courses, 4, end_row, len(courses_headers))
        if len(popular_courses) > 1:
            bar = BarChart()
            bar.type = "bar"
            bar.title = "المواد الأكثر طلباً"
            bar.style = 10
            bar.width = 20
            bar.height = 14
            data_ref = Reference(ws_courses, min_col=2, min_row=3, max_row=end_row)
            cats_ref = Reference(ws_courses, min_col=1, min_row=4, max_row=end_row)
            bar.add_data(data_ref, titles_from_data=True)
            bar.set_categories(cats_ref)
            ws_courses.add_chart(bar, f"A{end_row + 3}")
    _auto_adjust_columns(ws_courses)

    # ═══ حفظ وإرسال ═══
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    if orders_data or users_data:
        from datetime import datetime
        await update.message.reply_document(
            document=InputFile(
                output,
                filename=f"تقرير_البوت_{datetime.now().strftime('%Y-%m-%d')}.xlsx"),
            caption=(
                "✅ <b>تم تصدير التقرير الاحترافي بنجاح!</b>\n\n"
                "📄 يتضمن:\n"
                "• لوحة المعلومات (ملخص شامل)\n"
                "• سجل الطلبات (مع التخصص وتنسيق الحالات)\n"
                "• سجل المستخدمين\n"
                "• الإيرادات اليومية (مع رسم بياني)\n"
                "• الخدمات الشائعة (مع رسم دائري)\n"
                "• المواد الشائعة (مع رسم بياني)"
            ),
            parse_mode=ParseMode.HTML)
    else:
        await update.message.reply_text("📭 لا توجد بيانات للتصدير.")
